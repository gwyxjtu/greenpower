"""Generate the platform user manual as a Word document."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "绿电直连数据中心能源系统容量规划平台-使用说明.docx"
INK = RGBColor(0x2B, 0x24, 0x18)
ACCENT = RGBColor(0x8A, 0x6A, 0x2F)


def _set_run_font(run, size=12, bold=False, color=INK, east="宋体", west="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = west
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), west)
    rfonts.set(qn("w:hAnsi"), west)
    rfonts.set(qn("w:eastAsia"), east)
    rfonts.set(qn("w:cs"), west)


def _style_paragraph(p, before=0, after=6, line=1.35, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_title(doc, text):
    p = doc.add_paragraph()
    _style_paragraph(p, before=24, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    _set_run_font(run, size=22, bold=True, east="黑体")


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    _style_paragraph(p, before=0, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    _set_run_font(run, size=12, color=ACCENT, east="宋体")


def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    _style_paragraph(p, before=16 if level == 1 else 12, after=8)
    run = p.add_run(text)
    size = 16 if level == 1 else 13
    _set_run_font(run, size=size, bold=True, east="黑体")


def add_p(doc, text, first_line=True):
    p = doc.add_paragraph()
    _style_paragraph(p, before=0, after=6)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    _set_run_font(run, size=12)
    return p


def add_bullet(doc, text, numbered=False, n=None):
    p = doc.add_paragraph()
    _style_paragraph(p, before=0, after=3)
    p.paragraph_format.left_indent = Cm(0.74)
    prefix = f"{n}. " if numbered else "• "
    run = p.add_run(prefix + text)
    _set_run_font(run, size=12)


def add_note(doc, text):
    p = doc.add_paragraph()
    _style_paragraph(p, before=2, after=8)
    p.paragraph_format.left_indent = Cm(0.74)
    run = p.add_run("说明：" + text)
    _set_run_font(run, size=11, color=ACCENT)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        _set_run_font(run, size=11, bold=True, east="黑体")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        from docx.oxml import OxmlElement

        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E8DCC6")
        shd.set(qn("w:val"), "clear")
        tcpr.append(shd)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            _set_run_font(run, size=10.5)
    doc.add_paragraph()


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.8)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    styles = doc.styles["Normal"]
    styles.font.name = "Times New Roman"
    styles.font.size = Pt(12)
    styles.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    add_title(doc, "绿电直连数据中心能源系统")
    add_title(doc, "容量规划平台使用说明")
    add_subtitle(doc, f"面向业务人员操作手册　　{date.today().strftime('%Y年%m月%d日')}")

    add_h(doc, "一、平台是做什么的", 1)
    add_p(
        doc,
        "本平台用于智算中心绿电直连场景下的电源与接网容量规划。根据全年 8760 小时的负荷、风光资源和电价，自动给出风电、光伏、储能和接网变压器的推荐装机，并给出年化综合成本、度电成本、政策比例和逐时运行曲线。",
    )
    add_p(
        doc,
        "默认算例锚定宁夏 110 千伏两部制工商业电价口径。计算在后台进行，一次通常需要数分钟。多名用户可以同时登录；正在计算的任务最多同时 3 个，其余自动排队。",
    )
    add_p(doc, "访问地址：http://47.116.1.6:18501", first_line=True)

    add_h(doc, "二、登录", 1)
    add_bullet(doc, "用浏览器打开上述地址，进入登录页。", numbered=True, n=1)
    add_bullet(doc, "输入用户名和密码后点击「登录」。", numbered=True, n=2)
    add_bullet(doc, "登录后右上角可「退出登录」。", numbered=True, n=3)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["用户名", "cmcc"],
            ["密码", "CMCC123456"],
            ["同时登录", "支持，共用同一账号"],
            ["同时计算", "最多 3 个，超出的自动排队"],
        ],
    )
    add_note(doc, "页面上方会显示当前用户和排队情况。请勿把账号密码发到公开渠道。")

    add_h(doc, "三、推荐操作流程", 1)
    add_bullet(doc, "登录进入主页面。", numbered=True, n=1)
    add_bullet(doc, "如需从标准算例起步，点击「输入宁夏省案例参数」。", numbered=True, n=2)
    add_bullet(doc, "按项目情况修改运行政策、造价、电价，或上传本项目负荷与风光曲线。", numbered=True, n=3)
    add_bullet(doc, "点击「计算最优容量配置方案」，等待计算结束。", numbered=True, n=4)
    add_bullet(doc, "在「求解结果」中查看装机、成本、政策比例和曲线，必要时下载结果文件。", numbered=True, n=5)

    add_h(doc, "四、页面顶部按钮", 1)
    add_table(
        doc,
        ["按钮", "作用"],
        [
            [
                "输入宁夏省案例参数",
                "一键填入宁夏默认：接网变压器手动输入 60 MW，投资模式为负荷投资，光伏和风电 PPA 为 0，三项政策比例为 60% / 30% / 20%，并改回使用默认时序。",
            ],
            [
                "计算最优容量配置方案",
                "按当前页面参数启动计算。若已有 3 个任务在算，本任务会排队，页面提示排队位置。",
            ],
            ["退出登录", "退出当前浏览器会话，返回登录页。"],
        ],
    )

    add_h(doc, "五、运行与政策参数", 1)
    add_h(doc, "（一）接网变压器设计容量", 2)
    add_p(
        doc,
        "该项为选择，不是单独按钮。当前选中哪一项会直接显示在页面上：",
    )
    add_table(
        doc,
        ["选项", "含义"],
        [
            [
                "手动输入",
                "变压器容量按您填写的数值锁定。宁夏案例默认为 60 MW。此时模型只优化风电、光伏和储能。",
            ],
            [
                "由模型优化容量",
                "变压器也作为决策变量，在 0 到「接网变压器容量上限」（默认 200 MW）之间由模型一起选择。",
            ],
        ],
    )

    add_h(doc, "（二）投资模式", 2)
    add_table(
        doc,
        ["选项", "页面会出现", "成本怎么算"],
        [
            [
                "电源和专线由负荷投资",
                "资本回收系数、利率。光伏和风电 PPA 自动置为 0。",
                "电源与专线投资计入负荷侧年化综合成本。",
            ],
            [
                "电源和专线由发电企业投资",
                "光伏 PPA 电价、风电 PPA 电价。",
                "投资不计入负荷侧目标，负荷侧只计运行净支出；PPA 作为向发电企业支付的电价。PPA 必须低于网购电能量价。",
            ],
        ],
    )

    add_h(doc, "（三）政策比例", 2)
    add_p(doc, "三项均按百分数填写，范围为 0–100。")
    add_table(
        doc,
        ["名称", "含义", "宁夏默认"],
        [
            ["自发自用占总可用发电量比例下限", "本地消纳的风光电量，不得低于可用发电量的该比例。", "60%"],
            ["绿电发电量占总用电量比例下限", "年风光实际发电量不得低于年用电量的该比例。", "30%"],
            ["余电上网占总可用发电量比例上限", "余电上网不得超过可用发电量的该比例。", "20%"],
        ],
    )
    add_note(
        doc,
        "可用发电量按装机乘以年等效利用小时估算，不是实际发电量。实际发电允许弃风弃光，因此通常略小于可用发电量。「绿电发电量占比」看的是发电量相对用电量，不是负荷中本地绿电供上的份额。",
    )

    add_h(doc, "六、时间与负荷", 1)
    add_p(
        doc,
        "「负荷设计峰值」单位为 MW。选择「使用默认」时，平台按该峰值合成全年负荷，曲线约在峰值的一半到峰值之间，无需上传。",
    )
    add_p(
        doc,
        "选择「上传」时，必须提供全年 8760 小时负荷表。可先点「下载模板」。勾选「按设计峰值调整负荷」后，会把上传曲线的峰值拉到您填写的设计峰值，形状保持不变。",
    )

    add_h(doc, "七、设备造价与专线", 1)
    add_table(
        doc,
        ["参数", "单位", "宁夏默认"],
        [
            ["风电单位造价", "万元/MW", "410"],
            ["光伏单位造价", "万元/MW", "300"],
            ["储能单位造价", "万元/MWh", "80"],
            ["专线单位造价", "万元/km", "100"],
            ["专线距离", "km", "50"],
            ["光伏可用地上限", "亩", "1500（对应光伏装机大约不超过 137.5 MW）"],
            ["风电装机上限", "MW", "500"],
            ["储能能量上限", "MWh", "2000"],
            ["接网变压器容量上限", "MW", "200（仅「由模型优化容量」时作为上界）"],
        ],
    )

    add_h(doc, "八、电价", 1)
    add_p(doc, "页面按元/kWh 或元/kW·月填写，与日常电价口径一致。宁夏 110 千伏两部制默认如下。")
    add_table(
        doc,
        ["参数", "默认", "单位"],
        [
            ["所在电压等级容（需）量电价", "25.6", "元/kW·月"],
            ["所在电压等级输配电度电价", "0.060", "元/kWh"],
            ["系统运行费", "0.039", "元/kWh"],
            ["线损折价", "0.0071", "元/kWh"],
            ["政府性基金及附加", "0.0213", "元/kWh"],
            ["网购电能量价", "0.50", "元/kWh"],
            ["上网电价·峰 / 平 / 谷", "0.45 / 0.35 / 0.25", "元/kWh"],
            ["所在省份平均负荷率", "0.6", "无量纲"],
            ["年月份数", "12", "月"],
        ],
    )

    add_h(doc, "九、经济参数与储能", 1)
    add_table(
        doc,
        ["参数", "默认", "说明"],
        [
            ["项目设计使用年限", "15 年", "与利率一起决定资本回收系数。负荷投资模式下可直接改资本回收系数（默认约 0.1168）。"],
            ["利率", "8%", "仅负荷投资模式显示。"],
            ["充电 / 放电效率", "0.95", "大于 0、不超过 1。"],
            ["初始能量", "0 MWh", "年初储能能量。"],
            ["最大充电 / 放电功率", "50 MW", "储能功率上限，与储能能量容量不是同一个量。"],
        ],
    )

    add_h(doc, "十、新能源逐时出力", 1)
    add_p(
        doc,
        "光伏、风电均可选择「使用默认」或「上传」。默认曲线来自宁夏（银川）资源：光伏按实测出力系数，风电按年等效利用小时目标标定。",
    )
    add_p(
        doc,
        "选择上传时，必须提供 0 到 1 的出力系数，不要上传某座已建电站的兆瓦出力。可先下载模板。勾选「按目标小时数调整光伏出力」时，会把光伏曲线整体缩放到「所在省份光伏年等效利用小时数」。",
    )
    add_table(
        doc,
        ["文件", "表头（第 1 行）", "数据要求"],
        [
            ["负荷.csv", "小时,负荷（MW）", "8760 行，小时为 0–8759，负荷为 MW。"],
            ["光伏出力系数.csv", "小时,光伏出力系数", "8760 行，数值 0–1。"],
            ["风电出力系数.csv", "小时,风电出力系数", "8760 行，数值 0–1。"],
        ],
    )
    add_note(
        doc,
        "文件须为 UTF-8 逗号分隔。预览图会随上传即时更新，但只有点击「计算最优容量配置方案」后才会写入本次计算。",
    )

    add_h(doc, "十一、计算过程", 1)
    add_p(
        doc,
        "点击「计算最优容量配置方案」后，平台按全年 8760 小时建立优化模型并求解。页面会显示排队或计算状态。计算结束后自动展开结果。",
    )
    add_bullet(doc, "同时最多 3 个任务在计算，其余排队，互不影响各自参数。")
    add_bullet(doc, "默认求解间隙 1%、时限 600 秒。到时限若已有可行方案，结果仍可参考，但尚未证明全局最优。")
    add_bullet(doc, "若模型不可行，页面会提示放宽自发自用或绿电占比下限，或增大占地、变压器容量，并可下载不可行模型文件。")

    add_h(doc, "十二、如何阅读结果", 1)
    add_h(doc, "（一）状态与指标", 2)
    add_table(
        doc,
        ["栏目", "读法"],
        [
            ["已得到可接受最优解", "在设定间隙内认定最优，可直接采用装机方案。"],
            ["已到求解时限", "已有可行方案，装机可参考；请同时看「求解间隙」。间隙越大，离证明最优越远。"],
            ["年化综合成本", "目标值，越小越好，单位万元/年。"],
            ["度电成本", "年化综合成本除以年用电量。负荷投资时含投资年化；发电企业投资时仅为运行成本。"],
            ["最优装机", "风电（MW）、光伏（MW）、储能（MWh，能量容量）、接网变压器（MW）。"],
        ],
    )

    add_h(doc, "（二）成本拆解", 2)
    add_p(
        doc,
        "设备投资、专线投资列出的是一次性总价，计入年化综合成本时还要乘资本回收系数。容量电费、电量类支出、上网收益本身已是年值。不要把一次性投资直接与「万元/年」的目标值相加。",
    )

    add_h(doc, "（三）政策比例", 2)
    add_p(
        doc,
        "余电上网比例贴着上限，通常说明上网约束起作用。自发自用比例、绿电发电量占比应分别不低于您设定的下限。",
    )

    add_h(doc, "（四）逐时功率图", 2)
    add_p(
        doc,
        "全年曲线和「第几天」的日曲线，图例在图的左侧，做成可点选按钮。点击左侧名称即可显示或隐藏对应曲线（负荷、风电、光伏、购电、上网等）。日曲线还可查看充电、放电和荷电状态。",
    )
    add_p(
        doc,
        "平台会校核：各小时风电＋光伏－充电＋放电＋购电－上网应等于负荷。",
    )

    add_h(doc, "（五）下载", 2)
    add_table(
        doc,
        ["按钮", "内容"],
        [
            ["下载文字结果", "装机、成本拆解、政策比例、能量统计的文本。"],
            ["下载逐时数据", "8760 小时负荷、风光出力、储能充放、购电与上网。"],
            ["下载结构化结果", "供程序读取的结果摘要。"],
        ],
    )

    add_h(doc, "十三、常见问题", 1)
    add_bullet(doc, "点了变压器相关选项却看不出变化：请看「接网变压器设计容量」当前选中的是「手动输入」还是「由模型优化容量」。选手动输入时才会出现容量填写框。")
    add_bullet(doc, "计算很久：全年逐时优化规模大，数分钟属正常。可同时查看排队提示。")
    add_bullet(doc, "提示不可行：先放宽三项政策比例，或增大光伏用地、变压器容量上限。")
    add_bullet(doc, "提示 PPA 必须低于网购电能量价：把光伏、风电 PPA 调低，或提高网购电能量价。")
    add_bullet(doc, "上传失败：确认是 8760 行、UTF-8 逗号分隔，风光列为 0–1 的出力系数。闰年多出的 24 小时请不要直接使用。")
    add_bullet(doc, "外网打不开：先确认本机浏览器能打开访问地址；若仍失败，需在云服务器安全组放行 18501 端口。")

    add_h(doc, "附录  宁夏案例一键填入的内容", 1)
    add_table(
        doc,
        ["项目", "数值"],
        [
            ["接网变压器", "手动输入 60 MW"],
            ["投资模式", "电源和专线由负荷投资"],
            ["光伏 / 风电 PPA", "0 元/kWh"],
            ["资本回收系数", "0.1168（对应利率 8%、寿命 15 年）"],
            ["自发自用下限 / 绿电发电量下限 / 余电上网上限", "60% / 30% / 20%"],
            ["时序", "改回使用宁夏默认负荷与风光曲线"],
        ],
    )

    p = doc.add_paragraph()
    _style_paragraph(p, before=24, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run("— 完 —")
    _set_run_font(run, size=12, color=ACCENT)

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
